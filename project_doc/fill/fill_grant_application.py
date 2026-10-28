# -*- coding: utf-8 -*-
"""Заполняет шаблон заявки на грант Минобразования данными проекта Vuzhyk.

Шаблон: 12_124001_1_201178.doc (OLE) / .docx.
Чужие данные (Мацуганова, БГУИР, Казаков и т.д.) очищаются.
Неизвестные поля (дата рождения, специальность, научный руководитель,
суммы калькуляции) остаются пустыми.

Зависимости: pip install python-docx
Для .doc: установленный Microsoft Word (win32com).
"""
from __future__ import print_function

import os
import shutil
import sys

try:
    from docx import Document
except ImportError:
    sys.stderr.write("Нужен пакет python-docx: pip install python-docx\n")
    sys.exit(1)

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_DOCX = os.path.join(HERE, "12_124001_1_201178.docx")
TEMPLATE_DOC = os.path.join(HERE, "12_124001_1_201178.doc")
OUT_DOCX = os.path.join(HERE, "12_124001_1_201178_filled.docx")
OUT_DOC = os.path.join(HERE, "12_124001_1_201178_filled.doc")

# ---------------------------------------------------------------------------
# Данные проекта (только то, что есть в project_doc и в самом шаблоне ВГТУ)
# ---------------------------------------------------------------------------

APPLICANT_FULL = "Корнеев Павел Александрович"
APPLICANT_SHORT = "П.А. Корнеев"
PHONE = "+375 (25) 746-79-54"
FACULTY = "Факультет информационных технологий и робототехники"
UNIVERSITY_SHORT = "Витебский государственный технологический университет"
UNIVERSITY_FULL = (
    "Учреждение образования «Витебский государственный технологический университет»"
)
RECTOR_SHORT = "А.А. Кузнецов"
GROUP = "ИТ-13"
CATEGORY = "студент"

THEME = (
    "Разработка специализированной учебной среды программирования Vuzhyk "
    "для изучения языка Python в учреждениях общего среднего образования "
    "Республики Беларусь"
)

# Пункт 6 Положения (приказ Минобразования 24.06.2014 № 561)
DIRECTION = (
    "информационно-коммуникационные, авиационные и космические "
    "технологии и аппаратура"
)

JUSTIFICATION = {
    "1.": (
        "1. Цель научно-исследовательской работы, ее актуальность. "
        "Цель — разработать и научно обосновать специализированную "
        "интегрированную среду разработки Vuzhyk, адаптированную к учебной "
        "программе предмета «Информатика» учреждений общего среднего "
        "образования Республики Беларусь. "
        "Актуальность: в 2024/2025 учебном году изучение основ "
        "программирования в школах Беларуси перешло с языка Pascal на Python. "
        "Профессиональные IDE создают избыточную когнитивную нагрузку, не "
        "содержат исполнителя «Робот» в методике белорусских учебников, "
        "требуют отдельной установки интерпретатора и плохо работают на "
        "слабом школьном оборудовании. Необходима среда, которая совмещает "
        "дидактические средства (Черепаха с осями координат, встроенный "
        "задачник Робота, электронные учебники 6–7 классов, русскоязычная "
        "справка) с функциями профессиональной IDE."
    ),
    "2.": (
        "2. Научная идея. Образовательная IDE рассматривается как "
        "дидактический посредник между школьной программой и профессиональной "
        "практикой. Идея Vuzhyk — встроить в одну кроссплатформенную среду "
        "исполнителей белорусской программы информатики, задачник, справку и "
        "электронные учебники, сохранив привычные профессиональные приёмы "
        "(подсветка синтаксиса, автодополнение, отладка, запуск без сохранения). "
        "Название отсылает к белорусскому ужу в паре с языком Python."
    ),
    "3.": (
        "3. Структура и методы исследования. Анализ учебных программ и "
        "учебников 6–7 классов; сравнительный анализ существующих сред "
        "разработки; проектирование архитектуры на Qt; реализация встроенного "
        "интерпретатора Python, библиотек исполнителей Черепаха и Робот, "
        "редактора задач, справки; педагогическая апробация в учреждении "
        "общего среднего образования."
    ),
    "4.": (
        "4. Состояние рассматриваемой проблемы с учетом достижений "
        "современной науки. Профессиональные инструменты (IDLE, PyCharm и др.) "
        "не рассчитаны на урок: требуют отдельной установки Python, не "
        "визуализируют систему координат для Черепахи, не содержат задачника "
        "Робота из учебника 7 класса, имеют ошибки запуска поля Робота. "
        "Библиотеки Робота для Python существуют как разрозненные файлы. "
        "Без адекватного инструментария переход на Python снижает "
        "эффективность обучения и увеличивает нагрузку на учителя."
    ),
    "5.": (
        "5. Основные результаты исследования, их научная и практическая "
        "значимость, экономическая и социальная ценность. Создана "
        "функционирующая IDE Vuzhyk (Windows 7+, портативная и устанавливаемая "
        "сборки; готовится версия для Linux). Реализованы подсветка синтаксиса, "
        "автодополнение, справка и электронные учебники 6–7 классов, задачник "
        "и редактор задач Робота, Черепаха с координатной сеткой, запуск без "
        "сохранения, пониженное потребление оперативной памяти. Работа прошла "
        "университетский отборочный тур республиканского конкурса "
        "«100 идей для Беларуси». Имеется акт внедрения НИОК(Т)Р в ГУО "
        "«Средняя школа №2 г. Дубровно имени Ю. Смирнова» от 25.11.2025. "
        "Подготовлена статья на IX Международный научно-технический форум "
        "СТНО-2026. Практическая ценность — возможность вести урок "
        "информатики по белорусской программе без сторонних профессиональных "
        "IDE и на существующем парке школьных компьютеров."
    ),
    "6.": (
        "6. Возможные области использования результатов исследования. "
        "Уроки информатики 6–7 классов (далее — по мере выхода учебников "
        "следующих классов); внеурочная деятельность; подготовка учителей; "
        "учреждения образования с ограниченными вычислительными ресурсами. "
        "Сайт программы: vuzhyk.by. Репозиторий: github.com/Grossbeak/Vuzhyk."
    ),
    "7.": (
        "7. Сведения об участии соискателя гранта в научных исследованиях "
        "в данной области и принципиальное отличие данной "
        "научно-исследовательской работы. Соискатель является автором среды "
        "Vuzhyk, разработал архитектуру и основные модули (редактор, "
        "исполнители, справка, задачник). Принципиальное отличие: единая "
        "среда под белорусскую программу информатики, а не набор сторонних "
        "библиотек; встроенный интерпретатор; учёт слабого оборудования; "
        "подтверждённое школьное внедрение. Соавторы публикации по теме: "
        "А.С. Соколова, Д.В. Черненко."
    ),
}

CALENDAR = {
    # row index in table[1]: (name, result)
    1: (
        THEME,
        "Программный комплекс Vuzhyk, методика его применения на уроках "
        "информатики, материалы внедрения в учреждениях образования, "
        "заключительный отчет.",
    ),
    2: (
        "Теоретический этап. Анализ учебной программы информатики "
        "(включая требования 8 класса), сравнительное исследование "
        "образовательных IDE, формирование технического задания на развитие "
        "Vuzhyk.",
        "Аналитический обзор. Техническое задание.",
    ),
    3: (
        "Программная реализация развития Vuzhyk: синхронизация библиотек "
        "учитель–ученик по локальной сети, текстовые описания команд в "
        "подсказках, расширение справочной системы, доработка под учебную "
        "программу 8 класса.",
        "Прототип доработанной среды. Аннотированный отчет за 1-е полугодие.",
    ),
    4: (
        "Тестирование в учреждениях образования, исправление ошибок, "
        "подготовка дистрибутивов Windows и Linux, развитие сайта vuzhyk.by.",
        "Результаты испытаний. Доработанные дистрибутивы.",
    ),
    5: (
        "Внедрение, оформление документации и публикация результатов.",
        "Акты внедрения. Аннотированный отчет за год. Заключительный отчет.",
    ),
}

ARTICLE_TITLE = (
    "Разработка специализированной среды программирования для изучения "
    "Python в учреждениях общего среднего образования Республики Беларусь"
)
ARTICLE_BIBLIO = (
    "Корнеев П.А., Соколова А.С., Черненко Д.В. Разработка специализированной "
    "среды программирования для изучения Python в учреждениях общего среднего "
    "образования Республики Беларусь // Материалы IX Международного "
    "научно-технического форума «Современные технологии в науке и образовании» "
    "(СТНО-2026)."
)

# Чужие / шаблонные фрагменты → замена (пустая строка = удалить)
LEFTOVER_REPLACEMENTS = [
    ("Мацуганова Мария Дмитриевна", APPLICANT_FULL),
    ("М.Д. Мацуганова", APPLICANT_SHORT),
    ("М.Д.Мацуганова", APPLICANT_SHORT),
    ("В.Е. Казаков", ""),
    ("В.Е. Кзаков", ""),
    ("В.А. Богуш", RECTOR_SHORT),
    ("Н.Н. Федорович", ""),
    ("П.В. Камлач", ""),
    ("магистрант", CATEGORY),
    (
        "Разработка и оптимизация информационной системы анализа ассортимента предприятия",
        THEME,
    ),
    ("Ранее не получала", ""),
    ("27.05.2002", ""),
    ("+375333583524", PHONE),
    ('7-06-0612-03 «Системы управления информацией»', ""),
    ('7-06-0612-03 "Системы управления информацией"', ""),
    ("Информационные системы и технологии", ""),
    (
        "Междисциплинарные исследования. Перспективные зарождающиеся технологии.",
        DIRECTION,
    ),
    ("Дублируется название НИР", THEME),
    (
        "Общая сумма, далее разбивается на равные части по кварталам",
        "",
    ),
    ("Первый этап обязательно должен быть теоретическим", CALENDAR[2][0]),
]


def ensure_template_docx():
    if os.path.isfile(TEMPLATE_DOCX):
        return TEMPLATE_DOCX
    if not os.path.isfile(TEMPLATE_DOC):
        raise SystemExit("Нет шаблона: %s и %s" % (TEMPLATE_DOCX, TEMPLATE_DOC))
    try:
        import win32com.client  # type: ignore
    except ImportError:
        raise SystemExit(
            "Нет %s. Установите python-docx и Word либо сохраните шаблон как .docx."
            % TEMPLATE_DOCX
        )
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        doc = word.Documents.Open(os.path.abspath(TEMPLATE_DOC))
        wdFormatXMLDocument = 16
        doc.SaveAs2(os.path.abspath(TEMPLATE_DOCX), FileFormat=wdFormatXMLDocument)
        doc.Close(False)
    finally:
        word.Quit()
    return TEMPLATE_DOCX


def set_paragraph_text(paragraph, text):
    """Заменяет текст абзаца, сохраняя оформление первого run."""
    text = text if text is not None else ""
    runs = paragraph.runs
    if not runs:
        if text:
            paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def replace_in_text(text):
    if not text:
        return text
    new = text
    for old, repl in LEFTOVER_REPLACEMENTS:
        if old in new:
            new = new.replace(old, repl)
    # БГУИР в подписях калькуляции → ВГТУ
    new = new.replace(
        "«Белорусский государственный университет информатики и радиоэлектроники»",
        "«Витебский государственный технологический университет»",
    )
    new = new.replace(
        "«Белорусский государственный университет \nинформатики и радиоэлектроники»",
        "«Витебский государственный технологический университет»",
    )
    new = new.replace(
        "Белорусский государственный университет информатики и радиоэлектроники",
        UNIVERSITY_SHORT,
    )
    new = new.replace("НИЧ БГУИР", "НИЧ")
    new = new.replace(" БГУИР", "")
    new = new.replace("БГУИР", "")
    return new


def iter_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for table in doc.tables:
        for p in iter_table_paragraphs(table):
            yield p


def iter_table_paragraphs(table):
    seen = set()
    for row in table.rows:
        for cell in row.cells:
            el = cell._tc
            if id(el) in seen:
                continue
            seen.add(id(el))
            for p in cell.paragraphs:
                yield p
            for nested in cell.tables:
                for p in iter_table_paragraphs(nested):
                    yield p


def set_cell(table, row, col, text):
    cell = table.cell(row, col)
    lines = (text or "").split("\n")
    if not lines:
        lines = [""]
    while len(cell.paragraphs) < len(lines):
        cell.add_paragraph()
    for i, p in enumerate(cell.paragraphs):
        set_paragraph_text(p, lines[i] if i < len(lines) else "")


def save_as_doc(docx_path, doc_path):
    try:
        import win32com.client  # type: ignore
    except ImportError:
        print("pywin32 не установлен — .doc не создан, есть только .docx")
        return False
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        doc = word.Documents.Open(os.path.abspath(docx_path))
        wdFormatDocument = 0
        if os.path.isfile(doc_path):
            os.remove(doc_path)
        doc.SaveAs2(os.path.abspath(doc_path), FileFormat=wdFormatDocument)
        doc.Close(False)
    finally:
        word.Quit()
    return True


def fill(doc):
    # 1) Глобальная зачистка чужих данных во всех абзацах
    for p in iter_paragraphs(doc):
        original = p.text
        updated = replace_in_text(original)
        if updated != original:
            set_paragraph_text(p, updated)

    # 2) Заголовки и подписи
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("Научное направление:"):
            set_paragraph_text(
                p, "Научное направление: %s" % DIRECTION
            )
        elif t.startswith("Проведения научно-исследовательской работы"):
            set_paragraph_text(
                p,
                "Проведения научно-исследовательской работы «%s»" % THEME,
            )
        elif t == THEME or (t.startswith("«") and "Vuzhyk" in t):
            pass
        elif t.startswith("1.") and "Цель научно-исследовательской" in t:
            set_paragraph_text(p, JUSTIFICATION["1."])
        elif t.startswith("2.") and "Научная идея" in t:
            set_paragraph_text(p, JUSTIFICATION["2."])
        elif t.startswith("3.") and "Структура и методы" in t:
            set_paragraph_text(p, JUSTIFICATION["3."])
        elif t.startswith("4.") and "Состояние рассматриваемой" in t:
            set_paragraph_text(p, JUSTIFICATION["4."])
        elif t.startswith("5.") and "Основные результаты" in t:
            set_paragraph_text(p, JUSTIFICATION["5."])
        elif t.startswith("6.") and "Возможные области" in t:
            set_paragraph_text(p, JUSTIFICATION["6."])
        elif t.startswith("7.") and "Сведения об участии" in t:
            set_paragraph_text(p, JUSTIFICATION["7."])
        elif t.startswith("Список опубликованных работ студента"):
            set_paragraph_text(
                p,
                "Список опубликованных работ студента учреждения образования "
                "«Витебский государственный технологический университет»",
            )
        elif t.startswith("Автор") and "Фамилия" in t:
            set_paragraph_text(p, "Автор \t%s" % APPLICANT_SHORT)
        elif "Наименование научно-исследовательской работы" in t and t.startswith("("):
            pass
        elif set(t.replace("\t", "").replace(" ", "")) <= set("_"):
            # чисто подчёркивания под названием работы — подставить тему
            if len(t) > 20:
                set_paragraph_text(p, THEME)

    # Лишние строки-подчёркивания после пункта 7
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("________________") and "Vuzhyk" not in t:
            # не трогаем короткие линии подписей
            if len(t) > 40:
                set_paragraph_text(p, "")

    # Под названием калькуляции
    paras = list(doc.paragraphs)
    for i, p in enumerate(paras):
        if p.text.strip() == "затрат по выполнению научно-исследовательской работы":
            if i + 1 < len(paras):
                nxt = paras[i + 1].text.strip()
                if nxt.startswith("_") or "ассортимент" in nxt or not nxt:
                    set_paragraph_text(paras[i + 1], THEME)

    tables = doc.tables
    if len(tables) < 10:
        print("Ожидалось ≥10 таблиц, найдено %d" % len(tables))

    # TABLE[1] календарный план
    if len(tables) > 1:
        cal = tables[1]
        for row_idx, (name, result) in CALENDAR.items():
            if row_idx < len(cal.rows):
                set_cell(cal, row_idx, 1, name)
                set_cell(cal, row_idx, 3, "")  # стоимость неизвестна
                set_cell(cal, row_idx, 4, result)

    # TABLE[2] карточка заявки
    if len(tables) > 2:
        card = tables[2]
        set_cell(card, 0, 1, DIRECTION)
        set_cell(card, 1, 1, APPLICANT_FULL)
        set_cell(card, 2, 1, CATEGORY)
        set_cell(card, 3, 1, THEME)
        set_cell(card, 4, 1, "–")
        set_cell(card, 5, 1, UNIVERSITY_FULL)
        set_cell(card, 6, 1, "")  # шифр специальности неизвестен
        set_cell(card, 7, 1, "")  # прежние гранты неизвестны

    # TABLE[3] сведения о соискателе
    if len(tables) > 3:
        info = tables[3]
        set_cell(info, 0, 1, APPLICANT_FULL)
        set_cell(info, 1, 1, "")  # дата рождения неизвестна
        set_cell(
            info,
            2,
            1,
            "%s, группа %s" % (UNIVERSITY_SHORT, GROUP),
        )
        set_cell(info, 3, 1, "")  # дата поступления неизвестна
        set_cell(info, 4, 1, "")  # форма обучения неизвестна
        set_cell(info, 5, 1, "")  # специальность неизвестна
        set_cell(info, 6, 1, FACULTY)
        set_cell(info, 7, 1, "")  # кафедра неизвестна
        set_cell(info, 8, 1, PHONE)

    # TABLE[4]–[8] суммы — пустые (уже пустые в шаблоне)
    # TABLE[9] публикации
    if len(tables) > 9:
        pubs = tables[9]
        set_cell(pubs, 1, 1, ARTICLE_TITLE)
        set_cell(pubs, 1, 2, "статья")
        set_cell(pubs, 1, 3, ARTICLE_BIBLIO)
        set_cell(pubs, 2, 1, "")
        set_cell(pubs, 2, 2, "")
        set_cell(pubs, 2, 3, "")

    # Двухколоночные «сведения» (ФИО / место работы) — вписать соискателя
    for p in doc.paragraphs:
        raw = p.text
        if "отчество (полностью)" in raw and APPLICANT_FULL not in raw:
            set_paragraph_text(
                p, "отчество (полностью)\t%s\t" % APPLICANT_FULL
            )
        if raw.startswith("2.") and "Место работы" in raw:
            set_paragraph_text(
                p,
                "2. Место работы, должность\tстудент %s, группа %s\t"
                % (FACULTY.lower(), GROUP),
            )
        if raw.startswith("5.") and "Телефон (служебный)" in raw:
            set_paragraph_text(
                p, "5. Телефон (служебный)\t%s\t" % PHONE
            )


def main():
    src = ensure_template_docx()
    shutil.copyfile(src, OUT_DOCX)
    doc = Document(OUT_DOCX)
    fill(doc)
    doc.save(OUT_DOCX)
    print("Сохранено:", OUT_DOCX)
    try:
        if save_as_doc(OUT_DOCX, OUT_DOC):
            print("Сохранено:", OUT_DOC)
    except Exception as exc:
        print("Не удалось сохранить .doc:", exc)
        print("Откройте .docx и сохраните как .doc при необходимости.")


if __name__ == "__main__":
    main()
