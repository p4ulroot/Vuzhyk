#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Генерация полноценного доклада в Word (.docx) по подсистеме pyrob в Vuzhyk.
Требуется: pip install python-docx

Запуск из корня проекта:
    python tools/generate_pyrob_report_docx.py
или:
    python tools/generate_pyrob_report_docx.py --out "Доклад_Pyrob.docx"
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def set_cell_shading(cell, fill_hex: str) -> None:
    """Светло-серый фон ячейки (шестизначный RGB без #)."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill_hex)
    cell._tc.get_or_add_tcPr().append(shd)


def add_screenshot_placeholder(
    doc: Document,
    number: int,
    title: str,
    hint: str,
    height_cm: float = 7.5,
) -> None:
    """Блок под скриншот: подпись + рамка-ячейка с подсказкой."""
    p = doc.add_paragraph()
    run = p.add_run(f"Рисунок {number}. {title}")
    run.bold = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    row = table.rows[0]
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    row.height = Cm(height_cm)
    cell = row.cells[0]
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_shading(cell, "EEEEEE")
    inner = cell.paragraphs[0]
    inner.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = inner.add_run(f"\n{hint}\n\n")
    r.italic = True
    r.font.size = Pt(11)
    r2 = inner.add_run("(Удалите этот текст и вставьте изображение: «Вставка» → «Рисунки»)")
    r2.italic = True
    r2.font.size = Pt(10)
    doc.add_paragraph()  # отступ после блока


def set_document_defaults(doc: Document) -> None:
    """Поля, межстрочный интервал для основного текста."""
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)


def add_heading(doc: Document, text: str, level: int) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "Times New Roman"
        if level == 0:
            run.font.size = Pt(16)
        elif level == 1:
            run.font.size = Pt(14)
        else:
            run.font.size = Pt(14)


def add_para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph(text)
    if bold:
        for r in p.runs:
            r.bold = True


def add_title_page_table(doc: Document) -> None:
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Титульные данные (заполнить перед подачей)")
    r.bold = True

    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    rows = [
        ("Автор(ы)", "_________________________"),
        ("Организация / класс / команда", "_________________________"),
        ("Руководитель (при наличии)", "_________________________"),
        ("Город, год", "_________________________"),
        ("Конкурс / секция", "_________________________"),
    ]
    for i, (a, b) in enumerate(rows):
        table.rows[i].cells[0].text = a
        table.rows[i].cells[1].text = b
    doc.add_paragraph()


def build_document() -> Document:
    doc = Document()
    set_document_defaults(doc)

    # --- Титул ---
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(
        "АВТОМАТИЗАЦИЯ СОЗДАНИЯ УЧЕБНЫХ ЗАДАЧ ДЛЯ БИБЛИОТЕКИ PYROB\n"
        "В ИНТЕГРИРОВАННОЙ СРЕДЕ РАЗРАБОТКИ VUZHYK"
    )
    r.bold = True
    r.font.size = Pt(14)
    r.font.name = "Times New Roman"

    doc.add_paragraph()
    add_title_page_table(doc)

    # --- Аннотация ---
    add_heading(doc, "АННОТАЦИЯ", level=1)
    add_para(
        doc,
        "В работе рассматривается подсистема визуального конструирования и экспорта учебных заданий "
        "для робота-исполнителя на базе библиотеки pyrob в программном комплексе Vuzhyk (IDE на C++/Qt). "
        "Описана роль функции генерации исходного кода ProjectModel::generatePythonTask, преобразующей "
        "абстрактную модель поля (сетка, стены, старт, парковка, несколько вариантов уровня) в готовый модуль "
        "на языке Python, совместимый с ожидаемым интерфейсом задачи pyrob. Показаны принципы сериализации "
        "проекта в формат JSON, логика формирования методов load_level и check_solution, а также механизм "
        "однократного перемешивания порядка вариантов при загрузке модуля для исключения повторов в пределах "
        "цикла проверок. Отмечена практическая значимость для преподавания информатики и для сокращения доли "
        "ручного копирования шаблонного кода при составлении задач.",
    )
    add_para(doc, "Объём: ~3 стр.", bold=True)
    add_para(
        doc,
        "Ключевые слова: pyrob, визуальный редактор задач, генерация кода, Python, "
        "интегрированная среда разработки, обучение программированию.",
    )

    # --- Рисунок 1: после введения к интерфейсу ---
    add_heading(doc, "1. ВВЕДЕНИЕ", level=1)
    add_heading(doc, "1.1. Актуальность", level=2)
    add_para(
        doc,
        "Обучение программированию в школьной и вузовской практике широко опирается на задачи с исполнителем "
        "(робот на клетчатом поле). Библиотека pyrob предоставляет программный интерфейс для описания уровней, "
        "проверки решений и запуска сценариев. Ручное написание для каждой задачи вызовов API отнимает время, "
        "повышает риск ошибок и затрудняет сопровождение нескольких вариантов одного задания. Интеграция "
        "визуального редактора в IDE позволяет педагогу и учащемуся сосредоточиться на содержании задания.",
    )
    add_heading(doc, "1.2. Цель работы", level=2)
    add_para(
        doc,
        "Систематизировать описание механизма автоматического создания файла задачи pyrob (модуль .py) "
        "на основе данных редактора в проекте Vuzhyk.",
    )
    add_heading(doc, "1.3. Задачи", level=2)
    for item in (
        "Выделить состав модели данных задачи и формат сохранения проекта.",
        "Описать алгоритм генерации класса Task и взаимосвязь с библиотекой pyrob.",
        "Зафиксировать особенности поддержки нескольких вариантов уровня и проверки решения.",
        "Указать направления дальнейшего развития.",
    ):
        doc.add_paragraph(item, style="List Number")

    add_screenshot_placeholder(
        doc,
        1,
        "Интеграция редактора в IDE Vuzhyk",
        "Скриншот: главное окно приложения, меню «Сервис» / «Редактор задач pyrob» или вкладка редактора.",
        height_cm=6.5,
    )

    add_heading(doc, "2. ПОСТАНОВКА ЗАДАЧИ И СРЕДА РЕАЛИЗАЦИИ", level=1)
    add_para(
        doc,
        "Редактор задач pyrob встроен в приложение Vuzhyk и реализован с использованием Qt. Исходный код "
        "подсистемы расположен в каталоге src/pyrobeditor/. Пользователь задаёт геометрию поля, при необходимости "
        "несколько вариантов с различными размерами поля, сохраняет проект (.pyrob.json) и экспортирует задачу "
        "в виде файла Python (.py). Центральное преобразование «модель → текст программы» выполняет метод "
        "QString ProjectModel::generatePythonTask(const QString &taskId) const (файлы projectmodel.h, "
        "projectmodel.cpp). Запись на диск выполняется в PyrobEditorWidget.",
    )

    add_heading(doc, "3. МОДЕЛЬ ДАННЫХ И СОХРАНЕНИЕ ПРОЕКТА", level=1)
    add_heading(doc, "3.1. Структура варианта уровня", level=2)
    add_para(
        doc,
        "Для каждого варианта используется структура VariantState: размеры поля rows, cols; матрица типов клеток; "
        "матрицы внутренних стен; флаг и координаты старта; флаг и координаты парковки.",
    )
    add_heading(doc, "3.2. Формат проекта (JSON)", level=2)
    add_para(
        doc,
        "Методы toJson / fromJson формируют объект с полями variants, checks и массивом data с описаниями вариантов.",
    )

    add_screenshot_placeholder(
        doc,
        2,
        "Визуальный редактор задания pyrob",
        "Скриншот: сетка поля, панель инструментов, переключение вариантов, строки/столбцы.",
        height_cm=8.0,
    )

    add_heading(doc, "4. АЛГОРИТМ ГЕНЕРАЦИИ ФАЙЛА ЗАДАЧИ", level=1)
    add_para(
        doc,
        "Результат generatePythonTask — полный текст модуля: shebang, импорты, класс Task с методами load_level и "
        "check_solution. При нескольких вариантах в модуль добавляется _variant_order и random.shuffle при "
        "импорте. Для клеток «к заполнению» подключаются check_filled_cells и find_cells_to_be_filled из pyrob.tasks.",
    )
    add_para(
        doc,
        "Каталог по умолчанию для сохранения ориентирован на Python38/Lib/site-packages/pyrob/tasks относительно "
        "исполняемого файла IDE. Параметр taskId в текущей версии не встраивается в тело скрипта.",
    )

    add_screenshot_placeholder(
        doc,
        3,
        "Результат экспорта или сохранения задачи",
        "Скриншот: фрагмент сгенерированного .py (класс Task, load_level) или диалог «Сохранить задачу».",
        height_cm=7.0,
    )

    add_heading(doc, "5. РЕЗУЛЬТАТЫ И ПРАКТИЧЕСКАЯ ЗНАЧИМОСТЬ", level=1)
    for item in (
        "Сокращение ручного написания повторяющегося кода настройки поля и стен.",
        "Поддержка нескольких вариантов с разными размерами поля и перемешиванием порядка.",
        "Разделение проекта (JSON) и модуля задачи (PY) для удобных итераций.",
        "Единая IDE снижает порог входа по сравнению с внешним редактором и копированием кода.",
    ):
        doc.add_paragraph(item, style="List Number")

    add_heading(doc, "6. ЗАКЛЮЧЕНИЕ", level=1)
    add_para(
        doc,
        "Функция ProjectModel::generatePythonTask выполняет роль компилятора из доменной модели редактора в "
        "исполняемое описание Task, согласованное с библиотекой pyrob. Перспективы: использование taskId в "
        "комментариях, расширение проверок, валидация модели перед экспортом.",
    )

    add_heading(doc, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", level=1)
    for ref in (
        "Документация Python [Электронный ресурс]. URL: https://docs.python.org/3/ (дата обращения: _________).",
        "Материалы по библиотеке pyrob и каталогу tasks — по документации учебного комплекта.",
        "Qt Documentation [Электронный ресурс]. URL: https://doc.qt.io/ (дата обращения: _________).",
    ):
        doc.add_paragraph(ref, style="List Number")

    add_heading(doc, "ПРИЛОЖЕНИЕ", level=1)
    add_para(
        doc,
        "projectmodel.h, projectmodel.cpp — модель и генерация .py; "
        "PyrobEditorWidget.h, PyrobEditorWidget.cpp — интерфейс и сохранение.",
    )

    foot = doc.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = foot.add_run(
        "\nПеред подачей заполните титульный блок, даты в списке литературы и вставьте скриншоты вместо серых блоков."
    )
    fr.italic = True
    fr.font.size = Pt(11)

    return doc


def main() -> None:
    parser = argparse.ArgumentParser(description="Создать Word-доклад по pyrob (Vuzhyk).")
    parser.add_argument(
        "--out",
        type=Path,
        default=Path(__file__).resolve().parent.parent / "Доклад_Pyrob_Vuzhyk.docx",
        help="Путь к выходному .docx",
    )
    args = parser.parse_args()
    args.out.parent.mkdir(parents=True, exist_ok=True)

    doc = build_document()
    doc.save(str(args.out))
    print(f"Создан файл: {args.out.resolve()}")


if __name__ == "__main__":
    main()
